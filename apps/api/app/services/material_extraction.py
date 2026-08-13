from __future__ import annotations

import shutil
import subprocess
import tempfile
from math import sqrt
from pathlib import Path
from typing import Callable

import pypdfium2 as pdfium
import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader

from app.config import get_settings


MAX_EXTRACTED_CHARS = 120_000
DEFAULT_OCR_RENDER_SCALE = 2.2
FAST_OCR_RENDER_PIXELS = 800_000
MAX_OCR_RENDER_PIXELS = 2_000_000
MIN_USEFUL_OCR_CHARS = 40
TESSERACT_CONFIG = "--oem 1 --psm 6 -c tessedit_do_invert=0"
ExtractionProgress = Callable[[int, int], None]


def _ocr_image(image: Image.Image) -> str:
    prepared = image.convert("L")
    try:
        return pytesseract.image_to_string(
            prepared,
            lang="chi_sim",
            config=TESSERACT_CONFIG,
        ).strip()
    finally:
        if prepared is not image:
            prepared.close()


def _bounded_ocr_scale(width: float, height: float, max_pixels: int = MAX_OCR_RENDER_PIXELS) -> float:
    if width <= 0 or height <= 0:
        return DEFAULT_OCR_RENDER_SCALE
    return min(DEFAULT_OCR_RENDER_SCALE, sqrt(max_pixels / (width * height)))


def _ocr_resized(image: Image.Image, scale: float) -> str:
    if scale >= 1:
        return _ocr_image(image)
    resized = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))))
    try:
        return _ocr_image(resized)
    finally:
        resized.close()


def _ocr_pdf_page(pdf_page: pdfium.PdfPage, max_pixels: int) -> str:
    width, height = pdf_page.get_size()
    rendered = pdf_page.render(scale=_bounded_ocr_scale(width, height, max_pixels)).to_pil()
    try:
        return _ocr_image(rendered)
    finally:
        rendered.close()


def _adaptive_ocr_text(first_pass: str, retry: Callable[[], str]) -> str:
    if len(first_pass) >= MIN_USEFUL_OCR_CHARS:
        return first_pass
    retry_text = retry()
    return retry_text if len(retry_text) > len(first_pass) else first_pass


def _extract_pdf(path: Path, progress_callback: ExtractionProgress | None = None) -> str:
    reader = PdfReader(path)
    texts: list[str] = []
    weak_pages: list[int] = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        texts.append(text)
        if len(text) < MIN_USEFUL_OCR_CHARS:
            weak_pages.append(index)
    if weak_pages:
        document = pdfium.PdfDocument(str(path))
        try:
            total = len(weak_pages)
            for completed, index in enumerate(weak_pages, start=1):
                if progress_callback:
                    progress_callback(completed, total)
                pdf_page = document[index]
                try:
                    ocr_text = _adaptive_ocr_text(
                        _ocr_pdf_page(pdf_page, FAST_OCR_RENDER_PIXELS),
                        lambda current=pdf_page: _ocr_pdf_page(current, MAX_OCR_RENDER_PIXELS),
                    )
                finally:
                    pdf_page.close()
                if len(ocr_text) > len(texts[index]):
                    texts[index] = ocr_text
        finally:
            document.close()
    return "\n\n".join(f"--- 第{index + 1}页 ---\n{text}" for index, text in enumerate(texts))


def _extract_docx(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))
    return "\n".join(parts)


def _convert_office(path: Path, target_dir: Path) -> Path:
    settings = get_settings()
    executable = settings.soffice_path or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("Office材料转换服务不可用")
    profile = target_dir / "lo-profile"
    profile.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [
            executable,
            "--headless",
            "--nologo",
            "--norestore",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(target_dir),
            str(path),
        ],
        capture_output=True,
        timeout=120,
        check=False,
    )
    converted = target_dir / f"{path.stem}.pdf"
    if result.returncode != 0 or not converted.exists():
        raise RuntimeError("Office材料转换失败")
    return converted


def extract_material_text(path: Path, progress_callback: ExtractionProgress | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path, progress_callback)
    elif suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}:
        if progress_callback:
            progress_callback(1, 1)
        with Image.open(path) as image:
            width, height = image.size
            fast_scale = _bounded_ocr_scale(width, height, FAST_OCR_RENDER_PIXELS)
            if fast_scale < 1:
                fast = image.resize((max(1, int(width * fast_scale)), max(1, int(height * fast_scale))))
                retry_scale = _bounded_ocr_scale(width, height, MAX_OCR_RENDER_PIXELS)
                try:
                    text = _adaptive_ocr_text(
                        _ocr_image(fast),
                        lambda source=image, scale=retry_scale: _ocr_resized(source, scale),
                    )
                finally:
                    fast.close()
            else:
                text = _ocr_image(image)
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix in {".doc", ".xls", ".xlsx"}:
        with tempfile.TemporaryDirectory(prefix="material-") as temp:
            text = _extract_pdf(_convert_office(path, Path(temp)), progress_callback)
    else:
        raise RuntimeError(f"暂不支持读取 {suffix or '未知格式'} 材料")
    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not normalized:
        raise RuntimeError("材料中未识别到可读文字")
    return normalized[:MAX_EXTRACTED_CHARS]
