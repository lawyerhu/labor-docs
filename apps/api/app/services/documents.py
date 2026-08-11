from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import get_settings
from app.domain.calculations import calculate_claim
from app.domain.readiness import assess_readiness


PLACEHOLDER_COLOR = RGBColor(180, 45, 35)


@dataclass(frozen=True)
class GeneratedArtifact:
    kind: str
    filename: str
    path: Path


@dataclass(frozen=True)
class GenerationResult:
    readiness: str
    missing_fields: list[str]
    artifacts: list[GeneratedArtifact]


def _get(data: dict[str, Any], path: str, label: str) -> str:
    current: Any = data
    for part in path.split("."):
        if not isinstance(current, dict):
            current = None
            break
        current = current.get(part)
    if current in (None, "", []):
        return f"[待填入：{label}]"
    return str(current)


def _font(run, name: str = "宋体", size: float = 12, bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if "[待" in run.text:
        run.font.color.rgb = PLACEHOLDER_COLOR


def _configure_document(document: Document, landscape: bool = False) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.left_margin = Cm(1.8 if landscape else 3)
    section.right_margin = Cm(1.8 if landscape else 2.5)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    _font(run, size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    run = footer.add_run(" 页")
    _font(run, size=9)


def _paragraph(document: Document, text: str = "", *, indent: bool = True, bold_prefix: str | None = None, align=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    if align is not None:
        paragraph.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        left = paragraph.add_run(bold_prefix)
        _font(left, bold=True)
        right = paragraph.add_run(text[len(bold_prefix) :])
        _font(right)
    else:
        run = paragraph.add_run(text)
        _font(run)
    return paragraph


def _title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(14)
    run = paragraph.add_run(text)
    _font(run, "方正小标宋简体", 20, True)


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    _font(run, "黑体", 14, True)


def _party_text(data: dict[str, Any], role: str, prefix: str) -> list[str]:
    party = data.get("parties", {}).get(role, {})
    party_type = party.get("type") or ("company" if party.get("credit_code") else "individual")
    lines = [f"{prefix}：{party.get('name') or f'[待填入：{prefix}名称]'}。"]
    if party_type == "company":
        lines.append(f"住所地：{party.get('address') or f'[待填入：{prefix}住所地]'}。")
        lines.append(f"统一社会信用代码：{party.get('credit_code') or '[待填入：统一社会信用代码]'}。")
        lines.append(
            f"法定代表人：{party.get('legal_representative') or '[待填入：法定代表人姓名]'}，"
            f"职务：{party.get('legal_representative_title') or '[待填入：职务]'}。"
        )
    else:
        lines.append(
            f"性别：{party.get('gender') or '[待填入：性别]'}，出生日期：{party.get('birth_date') or '[待填入：出生日期]'}，"
            f"公民身份号码：{party.get('id_number') or '[待填入：身份证号码]'}。"
        )
        lines.append(f"住址：{party.get('address') or f'[待填入：{prefix}住所地]'}。")
    lines.append(f"联系方式：{party.get('contact') or '[待填入：联系电话]'}。")
    return lines


def _claim_lines(data: dict[str, Any], stage: str) -> list[str]:
    claims = data.get("claims") or []
    if not claims:
        return ["[待填入：仲裁请求]" if stage == "arbitration" else "[待填入：诉讼请求]"]
    lines: list[str] = []
    for claim in claims:
        title = claim.get("title") or "[待填入：请求内容]"
        amount = claim.get("amount")
        if amount in (None, ""):
            calculated = calculate_claim(claim.get("kind", "other"), claim.get("inputs") or {})
            amount_text = calculated.display
        else:
            amount_text = f"{float(amount):,.2f}元"
        basis = claim.get("basis") or "[待填入：金额、计算基数、期间或计算方式]"
        lines.append(f"{title}，金额为{amount_text}（计算依据：{basis}）。")
    if stage == "litigation":
        lines.append("本案诉讼费用由被告承担。")
    return lines


def _fact_text(data: dict[str, Any], stage: str) -> list[str]:
    employment = data.get("employment_facts") or {}
    facts = [
        f"申请人/原告于{employment.get('start_date') or '[待填入：入职日期]'}入职，"
        f"工作岗位为{employment.get('position') or '[待填入：工作岗位]'}，"
        f"工资标准为{employment.get('monthly_wage') or '[待填入：工资标准及构成]'}。",
        employment.get("summary") or "[待填入：劳动关系、工资支付、工作管理及解除经过等基本案情]",
    ]
    if stage == "litigation":
        arbitration = data.get("arbitration") or {}
        facts.append(
            f"本案经{arbitration.get('committee') or '[待填入：仲裁委员会]'}审理，作出"
            f"{arbitration.get('award_number') or '[待填入：仲裁裁决书案号]'}裁决，裁决结果为："
            f"{arbitration.get('result') or '[待填入：仲裁裁决结果]'}。"
        )
        facts.append(
            f"原告于{arbitration.get('service_date') or '[待填入：仲裁裁决送达日期]'}收到裁决，"
            f"履行情况为{arbitration.get('payment_status') or '[待填入：是否履行及已付款金额]'}。"
        )
    conflicts = data.get("unresolved_conflicts") or []
    facts.extend(f"[待核实：{item}]" for item in conflicts)
    legal = [item.get("citation") for item in data.get("legal_basis") or [] if item.get("verified")]
    facts.append("法律依据：" + "；".join(legal) if legal else "[待核验法律依据]")
    return facts


def _add_numbered(document: Document, lines: list[str]) -> None:
    for index, text in enumerate(lines, start=1):
        _paragraph(document, f"{index}. {text}", indent=False)


def _add_signature(document: Document, data: dict[str, Any], label: str, institution: str) -> None:
    _paragraph(document, "此致", indent=False)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(3)
    run = paragraph.add_run(institution)
    _font(run)
    signature = document.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signature.paragraph_format.space_before = Pt(2)
    signature.paragraph_format.space_after = Pt(0)
    signature.paragraph_format.line_spacing = 1.0
    run = signature.add_run(f"{label}：[待签名/盖章]　日期：[待填入：提交日期]")
    _font(run, size=11)


def _build_arbitration(path: Path, payload: dict[str, Any]) -> None:
    data = payload.get("data") or {}
    document = Document()
    _configure_document(document)
    _title(document, "劳动人事争议仲裁申请书")
    for line in _party_text(data, "initiating", "申请人") + _party_text(data, "opposing", "被申请人"):
        _paragraph(document, line, indent=False)
    _heading(document, "仲裁请求")
    _add_numbered(document, _claim_lines(data, "arbitration"))
    _heading(document, "事实与理由")
    for text in _fact_text(data, "arbitration"):
        _paragraph(document, text)
    committee = _get(data, "arbitration.committee", "劳动人事争议仲裁委员会")
    _add_signature(document, data, "申请人", committee)
    document.core_properties.title = "劳动人事争议仲裁申请书"
    document.save(path)


def _build_ordinary_complaint(path: Path, payload: dict[str, Any]) -> None:
    data = payload.get("data") or {}
    document = Document()
    _configure_document(document)
    _title(document, "民事起诉状")
    for line in _party_text(data, "initiating", "原告") + _party_text(data, "opposing", "被告"):
        _paragraph(document, line, indent=False)
    _heading(document, "诉讼请求")
    _add_numbered(document, _claim_lines(data, "litigation"))
    _heading(document, "事实与理由")
    for text in _fact_text(data, "litigation"):
        _paragraph(document, text)
    court = _get(data, "court", "管辖人民法院")
    _add_signature(document, data, "起诉人", court)
    document.core_properties.title = "劳动争议民事起诉状（普通式）"
    document.save(path)


def _set_cell(cell, value: str, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(value)
    _font(run, size=10, bold=bold)


def _set_table_geometry(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dxa = [round(width / 2.54 * 1440) for width in widths_cm]
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(dxa)))
    table_width.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, dxa):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _build_element_complaint(path: Path, payload: dict[str, Any]) -> None:
    data = payload.get("data") or {}
    document = Document()
    _configure_document(document)
    _title(document, "民事起诉状（劳动争议要素式）")
    rows = [
        ("一、当事人信息", ""),
        ("原告", "\n".join(_party_text(data, "initiating", "原告"))),
        ("被告", "\n".join(_party_text(data, "opposing", "被告"))),
        ("二、诉讼请求", "\n".join(f"{i}. {line}" for i, line in enumerate(_claim_lines(data, "litigation"), 1))),
        ("三、劳动关系要素", "\n".join(_fact_text(data, "litigation")[:2])),
        ("四、仲裁前置情况", "\n".join(_fact_text(data, "litigation")[2:4])),
        ("五、事实、理由及依据", "\n".join(_fact_text(data, "litigation")[4:] or ["[待填入：事实、理由及依据]"])),
        ("六、受诉法院", _get(data, "court", "管辖人民法院")),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    _set_table_geometry(table, [4, 11.5])
    for index, (label, value) in enumerate(rows):
        _set_cell(table.cell(index, 0), label, bold=True)
        _set_cell(table.cell(index, 1), value)
        table.cell(index, 0).width = Cm(4)
        table.cell(index, 1).width = Cm(11.5)
    _add_signature(document, data, "起诉人", _get(data, "court", "管辖人民法院"))
    document.core_properties.title = "劳动争议民事起诉状（要素式）"
    document.save(path)


def _source_pdf(path: Path, temp_dir: Path) -> Path:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        PdfReader(path)
        return path
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        image = Image.open(path).convert("RGB")
        target = temp_dir / f"{path.stem}.pdf"
        image.save(target, "PDF", resolution=150)
        return target
    if suffix in {".doc", ".docx", ".xls", ".xlsx"}:
        settings = get_settings()
        soffice = settings.soffice_path or shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise RuntimeError(f"无法转换Office文件：{path.name}，请配置 SOFFICE_PATH")
        profile = temp_dir / "libreoffice-profile"
        profile.mkdir(parents=True, exist_ok=True)
        subprocess.run(
            [
                soffice,
                "--headless",
                "--nologo",
                "--norestore",
                "--nolockcheck",
                "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                str(temp_dir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
        target = temp_dir / f"{path.stem}.pdf"
        if not target.exists():
            raise RuntimeError(f"Office文件转换失败：{path.name}")
        return target
    raise RuntimeError(f"不支持的证据格式：{path.suffix}")


def _footer_overlay(width: float, height: float, page: int, total: int) -> Any:
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except KeyError:
        pass
    stream = io.BytesIO()
    overlay = canvas.Canvas(stream, pagesize=(width, height))
    overlay.setFont("STSong-Light", 9)
    overlay.drawCentredString(width / 2, 16, f"第{page}/{total}页")
    overlay.save()
    stream.seek(0)
    return PdfReader(stream).pages[0]


def _assemble_evidence(evidence: list[dict[str, Any]], path: Path) -> list[str]:
    with tempfile.TemporaryDirectory(prefix="labor-docs-") as temporary:
        temp_dir = Path(temporary)
        sources: list[tuple[dict[str, Any], PdfReader]] = []
        for item in evidence:
            source = _source_pdf(Path(item["stored_path"]), temp_dir)
            sources.append((item, PdfReader(source)))
        total = sum(len(reader.pages) for _, reader in sources)
        writer = PdfWriter()
        page_ranges: list[str] = []
        current = 1
        bookmark_starts: list[tuple[str, int]] = []
        for index, (item, reader) in enumerate(sources, start=1):
            start = current
            bookmark_starts.append((f"证据{index}：{item.get('name') or item.get('original_name')}", len(writer.pages)))
            for source_page in reader.pages:
                source_page.merge_page(_footer_overlay(float(source_page.mediabox.width), float(source_page.mediabox.height), current, total))
                writer.add_page(source_page)
                current += 1
            end = current - 1
            page_ranges.append(str(start) if start == end else f"{start}—{end}")
        for title, page_index in bookmark_starts:
            writer.add_outline_item(title, page_index)
        writer.add_metadata({"/Title": "证据材料", "/Author": ""})
        with path.open("wb") as output:
            writer.write(output)
    return page_ranges


def _build_catalog(path: Path, evidence: list[dict[str, Any]], page_ranges: list[str]) -> None:
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except KeyError:
        pass
    document = SimpleDocTemplate(
        str(path), pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, title="证据目录",
    )
    normal = ParagraphStyle("catalog", fontName="STSong-Light", fontSize=9, leading=13)
    centered = ParagraphStyle("catalog-center", parent=normal, alignment=TA_CENTER)
    title = ParagraphStyle("catalog-title", parent=centered, fontSize=18, leading=24, spaceAfter=12)
    headers = ["证据编号", "证据名称", "来源", "证明目的", "页码"]
    rows: list[list[Any]] = [[Paragraph(value, centered) for value in headers]]
    for row_index, item in enumerate(evidence, start=1):
        values = [
            str(row_index),
            item.get("name") or item.get("original_name") or f"证据{row_index}",
            item.get("source") or "[待填入：来源]",
            item.get("purpose") or "[待填入：证明目的]",
            page_ranges[row_index - 1],
        ]
        rows.append([Paragraph(str(value), centered if index in {0, 4} else normal) for index, value in enumerate(values)])
    table = Table(rows, colWidths=[2.2 * cm, 5 * cm, 4.2 * cm, 11.5 * cm, 2.5 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF6")),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#64748B")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    document.build([Paragraph("证据目录", title), Spacer(1, 4), table])


def build_case_package(
    *,
    case_id: str,
    payload: dict[str, Any],
    evidence_items: list[dict[str, Any]],
    output_dir: Path,
) -> GenerationResult:
    assessment = assess_readiness(payload)
    if not assessment.can_generate:
        raise ValueError("必须先选择案件阶段和申请人/原告一方")
    case_dir = output_dir / case_id
    case_dir.mkdir(parents=True, exist_ok=True)
    artifacts: list[GeneratedArtifact] = []
    stage = payload["case_stage"]
    if stage == "arbitration":
        target = case_dir / "01-劳动人事争议仲裁申请书.docx"
        _build_arbitration(target, payload)
        artifacts.append(GeneratedArtifact("pleading", target.name, target))
    else:
        ordinary = case_dir / "01-民事起诉状.docx"
        _build_ordinary_complaint(ordinary, payload)
        artifacts.append(GeneratedArtifact("pleading", ordinary.name, ordinary))

    page_ranges: list[str] = []
    if evidence_items:
        evidence_path = case_dir / "03-证据材料.pdf"
        page_ranges = _assemble_evidence(evidence_items, evidence_path)
    catalog = case_dir / "02-证据目录.pdf"
    _build_catalog(catalog, evidence_items, page_ranges)
    artifacts.append(GeneratedArtifact("evidence_catalog", catalog.name, catalog))
    if evidence_items:
        artifacts.append(GeneratedArtifact("evidence_materials", "03-证据材料.pdf", case_dir / "03-证据材料.pdf"))
    return GenerationResult(assessment.readiness, assessment.missing_fields, artifacts)
