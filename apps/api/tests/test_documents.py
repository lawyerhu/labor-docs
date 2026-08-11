from pathlib import Path
from types import SimpleNamespace

from PIL import Image
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.services.documents import _source_pdf, build_case_package


def _all_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def _make_pdf(path: Path, label: str, pages: int = 1) -> None:
    output = canvas.Canvas(str(path))
    for number in range(1, pages + 1):
        output.drawString(72, 760, f"{label} page {number}")
        output.showPage()
    output.save()


def test_litigation_package_with_missing_facts_contains_placeholders(tmp_path: Path):
    result = build_case_package(
        case_id="case-empty",
        payload={"case_stage": "litigation", "party_side": "employer", "data": {}},
        evidence_items=[],
        output_dir=tmp_path,
    )

    names = {artifact.filename for artifact in result.artifacts}
    assert names == {"01-民事起诉状.docx", "02-证据目录.pdf"}
    assert "03-证据材料.pdf" not in names
    ordinary = tmp_path / "case-empty" / "01-民事起诉状.docx"
    assert "[待填入：原告名称]" in _all_text(ordinary)
    catalog = PdfReader(tmp_path / "case-empty" / "02-证据目录.pdf")
    assert len(catalog.pages) == 1
    assert float(catalog.pages[0].mediabox.width) > float(catalog.pages[0].mediabox.height)


def test_evidence_package_uses_actual_files_and_continuous_pages(tmp_path: Path):
    source_a = tmp_path / "a.pdf"
    source_b = tmp_path / "b.pdf"
    _make_pdf(source_a, "A", 2)
    _make_pdf(source_b, "B", 1)

    result = build_case_package(
        case_id="case-evidence",
        payload={"case_stage": "arbitration", "party_side": "worker", "data": {}},
        evidence_items=[
            {"id": "a", "name": "劳动合同", "source": "双方签署", "purpose": "证明劳动关系。", "stored_path": str(source_a)},
            {"id": "b", "name": "工资记录", "source": "银行", "purpose": "证明工资标准。", "stored_path": str(source_b)},
        ],
        output_dir=tmp_path,
    )

    evidence = tmp_path / "case-evidence" / "03-证据材料.pdf"
    assert evidence.exists()
    reader = PdfReader(evidence)
    assert len(reader.pages) == 3
    assert len(reader.outline) == 2
    catalog = PdfReader(tmp_path / "case-evidence" / "02-证据目录.pdf")
    assert len(catalog.pages) == 1
    assert float(catalog.pages[0].mediabox.width) > float(catalog.pages[0].mediabox.height)
    assert result.readiness == "formal_with_placeholders"


def test_evidence_package_converts_images_to_continuous_pdf(tmp_path: Path):
    image = tmp_path / "聊天截图.png"
    Image.new("RGB", (400, 200), "white").save(image)

    result = build_case_package(
        case_id="case-image-evidence",
        payload={"case_stage": "arbitration", "party_side": "worker", "data": {}},
        evidence_items=[
            {
                "id": "image",
                "name": "聊天截图",
                "source": "当事人提供",
                "purpose": "证明沟通内容。",
                "stored_path": str(image),
            }
        ],
        output_dir=tmp_path,
    )

    evidence = tmp_path / "case-image-evidence" / "03-证据材料.pdf"
    assert evidence.exists()
    reader = PdfReader(evidence)
    assert len(reader.pages) == 1
    assert "第1/1页" in (reader.pages[0].extract_text() or "")
    assert any(artifact.filename == "03-证据材料.pdf" for artifact in result.artifacts)


def test_office_conversion_uses_an_isolated_libreoffice_profile(monkeypatch, tmp_path: Path):
    source = tmp_path / "材料.docx"
    source.write_bytes(b"test office input")
    calls = []

    def fake_run(args, **kwargs):
        calls.append((args, kwargs))
        (tmp_path / "材料.pdf").write_bytes(b"%PDF-1.4")

    monkeypatch.setattr("app.services.documents.get_settings", lambda: SimpleNamespace(soffice_path="soffice"))
    monkeypatch.setattr("app.services.documents.subprocess.run", fake_run)

    converted = _source_pdf(source, tmp_path)

    assert converted == tmp_path / "材料.pdf"
    command = calls[0][0]
    assert any(argument.startswith("-env:UserInstallation=file:///") for argument in command)
    assert "--norestore" in command
    assert "--nolockcheck" in command
